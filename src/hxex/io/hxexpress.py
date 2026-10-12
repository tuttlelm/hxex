# SPDX-License-Identifier: GPL-3.0-or-later
"""HX-Express style data: one spreadsheet (.xlsx or .csv) per sample, peptide and charge.

Row 1 holds one time label per spectrum (``undeut``, ``TD``, ``60 s``, ``5 min``); each label
sits above an (m/z, intensity) column pair and the data start in row 2.
"""

from __future__ import annotations

import csv
import logging
from collections import Counter
from collections.abc import Sequence
from pathlib import Path

from ..spectra import PeptideEntry, Spectrum
from ._common import clean_spectrum, to_number
from .entries import filter_entries, sort_entries
from .errors import ReadError
from .times import (
    DEFAULT_FULLDEUT_TIME,
    DEFAULT_UNDEUT_TIME,
    is_control_label,
    parse_hxexpress_time,
)

log = logging.getLogger(__name__)


def parse_hxexpress_filename(name: str) -> PeptideEntry:
    """Entry from ``SAMPLE_START-END-PEPTIDE-zCHARGE-usertext.xlsx``.

    ``SAMPLE`` may contain underscores and hyphen-free text; the text after the last
    underscore of the first field is the start residue.
    """
    labels = name.split("-")
    try:
        if len(labels) < 5:
            raise ValueError("expected at least five '-' separated fields")
        sample, start = labels[0].rsplit("_", 1)
        end = labels[-4]
        charge = labels[-2]
        if not charge.startswith("z"):
            raise ValueError(f"charge field {charge!r} does not start with 'z'")
        return PeptideEntry(
            file=name,
            sample=sample,
            peptide=labels[-3],
            peptide_range=f"{start}-{end}",
            start_seq=int(start),
            end_seq=int(end),
            charge=int(charge[1:]),
        )
    except ValueError as exc:
        raise ReadError(
            f"cannot read sample/peptide/charge from file name {name!r} "
            f"(expected SAMPLE_START-END-PEPTIDE-zCHARGE-text.xlsx): {exc}"
        ) from exc


def scan_hxexpress(
    data_dir: str | Path,
    *,
    samples: Sequence[str] | None = None,
    peptide_ranges: Sequence[str] | None = None,
) -> list[PeptideEntry]:
    """List the entries for the ``.xlsx`` files in ``data_dir`` named in the HX-Express way.

    Files whose names do not parse are skipped with a warning.
    """
    entries = []
    for path in sorted(Path(data_dir).glob("*.xlsx")):
        if path.name.startswith("~$"):  # Excel lock file
            continue
        try:
            entries.append(parse_hxexpress_filename(path.name))
        except ReadError as exc:
            log.warning("skipping %s: %s", path.name, exc)
    return sort_entries(filter_entries(entries, samples=samples, peptide_ranges=peptide_ranges))


def _read_rows(path: Path) -> list[tuple]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            finally:
                wb.close()
        if suffix == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as fh:
                return [tuple(r) for r in csv.reader(fh)]
    except OSError as exc:
        raise ReadError(f"cannot read {path}: {exc}") from exc
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise ReadError(f"cannot read {path}: {exc}") from exc
    raise ReadError(f"{path}: unsupported file type {suffix!r} (use .xlsx or .csv)")


def read_hxexpress_file(
    path: str | Path,
    entry: PeptideEntry,
    *,
    data_id: int = 0,
    fulldeut_time: float = DEFAULT_FULLDEUT_TIME,
    undeut_time: float = DEFAULT_UNDEUT_TIME,
) -> list[Spectrum]:
    """Read every spectrum in one HX-Express file.

    Spectra come in column order. ``rep`` counts the spectra with the same time so far. A
    control label and a time label that give the same time are logged as a warning (they cannot
    be told apart afterwards; set ``undeut_time`` or ``fulldeut_time`` to another value).
    """
    path = Path(path)
    rows = _read_rows(path)
    if not rows:
        raise ReadError(f"{path}: file is empty")
    header = rows[0]
    labelled = [
        (col, str(cell).strip())
        for col, cell in enumerate(header)
        if cell is not None and str(cell).strip() not in ("", "nan")
    ]
    if not labelled:
        raise ReadError(f"{path}: first row has no time labels")
    seen: Counter[float] = Counter()
    _warn_control_clash(path, [(label, parse_hxexpress_time(label, fulldeut_time, undeut_time))
                               for _, label in labelled if _parses(label)])
    spectra = []
    for col, label in labelled:
        try:
            time = parse_hxexpress_time(label, fulldeut_time, undeut_time)
        except ReadError as exc:
            raise ReadError(f"{path}, column {col + 1}: {exc}") from exc
        seen[time] += 1
        mz_col = [r[col] if col < len(r) else None for r in rows[1:]]
        in_col = [r[col + 1] if col + 1 < len(r) else None for r in rows[1:]]
        where = f"{path.name}, label {label!r}"
        mz, intensity = clean_spectrum(
            [to_number(c, where) for c in mz_col], [to_number(c, where) for c in in_col]
        )
        spectra.append(
            Spectrum(
                mz, intensity, entry, time, seen[time],
                data_id=data_id, label=label, spectrum_id=f"{entry.file}#{col + 1}",
            )
        )
    return spectra


def _parses(label: str) -> bool:
    try:
        parse_hxexpress_time(label)
    except ReadError:
        return False
    return True


def _warn_control_clash(where, labelled_times) -> None:
    """Warn when a control label and a time label give the same time (shared with SpecExport)."""
    control = {t for label, t in labelled_times if is_control_label(label)}
    timed = {t for label, t in labelled_times if not is_control_label(label)}
    clash = sorted(control & timed)
    if clash:
        log.warning(
            "%s: a control and a time point both have time %s, so they cannot be told apart; "
            "set undeut_time or fulldeut_time to a value no time point has (for example -1)",
            where, ", ".join(f"{t:g}" for t in clash),
        )
